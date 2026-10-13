"""
Étape 1 — Ingestion des documents.

Rôle :
- Charger tous les documents présents dans data/docs/ (.txt, .md, .pdf, .pptx)
- Les découper en chunks (morceaux de texte) adaptés à la recherche sémantique

Ce module est importé par build_index.py (étape 2). Il peut aussi être lancé
directement pour vérifier rapidement combien de chunks seraient générés.
"""

import os
import re
from pathlib import Path

from langchain_community.document_loaders import TextLoader, PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "docs"

# Taille des chunks : à ajuster selon les résultats obtenus (voir README).
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 180

# En-têtes de chapitre répétés sur chaque page d'un PDF (ex: "CHAPITRE 2. CADRE DU
# SERVICE...") : du bruit qui pollue les chunks et que le LLM a tendance à recopier
# tel quel dans ses réponses. On les retire à l'extraction.
CHAPTER_HEADER_RE = re.compile(r"(?im)^\s*CHAPITRE\s+\d+\.\s*.*$")


def clean_text(text: str) -> str:
    """Retire le bruit récurrent de l'extraction PDF (en-têtes de chapitre répétés,
    lignes vides en excès) sans toucher au contenu utile."""
    text = CHAPTER_HEADER_RE.sub("", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pptx(file_path: Path) -> list[Document]:
    """Charge un fichier PowerPoint (.pptx), une diapositive = un Document.
    Extrait le texte des zones de texte et des tableaux ; ignore les images."""
    from pptx import Presentation

    prs = Presentation(str(file_path))
    documents = []

    for i, slide in enumerate(prs.slides, start=1):
        morceaux = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraphe in shape.text_frame.paragraphs:
                    ligne = "".join(run.text for run in paragraphe.runs).strip()
                    if ligne:
                        morceaux.append(ligne)
            if shape.has_table:
                for row in shape.table.rows:
                    ligne_tableau = " | ".join(cell.text.strip() for cell in row.cells)
                    if ligne_tableau.strip(" |"):
                        morceaux.append(ligne_tableau)

        texte_slide = "\n".join(morceaux).strip()
        if texte_slide:
            documents.append(
                Document(
                    page_content=texte_slide,
                    metadata={"source": str(file_path), "slide": i},
                )
            )

    return documents


def load_documents(docs_dir: Path = DATA_DIR) -> list[Document]:
    """Charge tous les fichiers supportés (.txt, .md, .pdf, .pptx) d'un dossier."""
    documents: list[Document] = []

    if not docs_dir.exists():
        raise FileNotFoundError(f"Le dossier {docs_dir} n'existe pas.")

    files = sorted(docs_dir.rglob("*"))
    if not files:
        print(f"[ingest] Attention : aucun fichier trouvé dans {docs_dir}")

    for file_path in files:
        if file_path.is_dir():
            continue

        suffix = file_path.suffix.lower()
        try:
            if suffix in (".txt", ".md"):
                loader = TextLoader(str(file_path), encoding="utf-8")
                loaded = loader.load()
            elif suffix == ".pdf":
                loader = PyPDFLoader(str(file_path))
                loaded = loader.load()
            elif suffix == ".pptx":
                loaded = load_pptx(file_path)
                if not loaded:
                    print(f"[ingest] Attention : aucun texte extrait de {file_path.name}")
            else:
                print(f"[ingest] Type de fichier ignoré : {file_path.name}")
                continue

            for doc in loaded:
                doc.page_content = clean_text(doc.page_content)
            documents.extend(loaded)

            print(f"[ingest] Chargé : {file_path.name} ({len(loaded)} section(s))")
        except Exception as exc:
            print(f"[ingest] Erreur lors du chargement de {file_path.name} : {exc}")

    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    """Découpe les documents en chunks avec chevauchement (overlap)."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n## ", "\n# ", "\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    return chunks


def load_and_split(docs_dir: Path = DATA_DIR) -> list[Document]:
    documents = load_documents(docs_dir)
    chunks = split_documents(documents)
    return chunks


if __name__ == "__main__":
    chunks = load_and_split()
    print(f"\n[ingest] Total : {len(chunks)} chunks générés.")
    if chunks:
        print("\n--- Exemple de chunk ---")
        print(chunks[0].page_content[:400])
        print("------------------------")
